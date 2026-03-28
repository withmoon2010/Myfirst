from __future__ import annotations

import logging
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from src.processing.models import Newsletter, SummarizedItem, SpecialTopic

logger = logging.getLogger(__name__)


def _set_cell_shading(cell, color_hex: str) -> None:
    """Set background color on a table cell."""
    shading = cell._element.get_or_add_tcPr()
    shading_elem = shading.makeelement(
        qn("w:shd"),
        {
            qn("w:fill"): color_hex,
            qn("w:val"): "clear",
        },
    )
    shading.append(shading_elem)


class DocxFormatter:
    def __init__(self) -> None:
        pass

    def render(self, newsletter: Newsletter) -> Document:
        doc = Document()

        # ── Page setup ──
        section = doc.sections[0]
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

        # ── Styles setup ──
        style = doc.styles["Normal"]
        font = style.font
        font.name = "Calibri"
        font.size = Pt(11)
        font.color.rgb = RGBColor(0x33, 0x41, 0x55)

        # ── Title ──
        title = doc.add_heading("DAILY MACRO BRIEF", level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.runs[0]
        run.font.size = Pt(24)
        run.font.color.rgb = RGBColor(0x1A, 0x23, 0x32)

        date_para = doc.add_paragraph()
        date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        date_run = date_para.add_run(newsletter.date.strftime("%A, %B %d, %Y"))
        date_run.font.size = Pt(12)
        date_run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

        doc.add_paragraph()  # spacer

        # ── Key Takeaways ──
        if newsletter.key_takeaways:
            h = doc.add_heading("KEY TAKEAWAYS", level=1)
            h.runs[0].font.color.rgb = RGBColor(0x1A, 0x23, 0x32)

            for i, takeaway in enumerate(newsletter.key_takeaways, 1):
                para = doc.add_paragraph()
                num_run = para.add_run(f"{i}. ")
                num_run.bold = True
                num_run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
                text_run = para.add_run(takeaway)
                text_run.font.size = Pt(11)
                para.paragraph_format.space_after = Pt(6)

            doc.add_paragraph()  # spacer

        # ── Risk Alerts ──
        if newsletter.risk_alerts:
            h = doc.add_heading("RISK ALERTS", level=1)
            h.runs[0].font.color.rgb = RGBColor(0xDC, 0x26, 0x26)

            table = doc.add_table(rows=1, cols=3)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.style = "Table Grid"

            # Header row
            headers = ["Alert", "Summary", "Market Implication"]
            header_row = table.rows[0]
            for i, header_text in enumerate(headers):
                cell = header_row.cells[i]
                cell.text = header_text
                _set_cell_shading(cell, "DC2626")
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                        run.font.size = Pt(10)

            for alert in newsletter.risk_alerts:
                row = table.add_row()
                row.cells[0].text = alert.title
                row.cells[1].text = alert.summary
                row.cells[2].text = alert.market_implication
                for cell in row.cells:
                    _set_cell_shading(cell, "FEF2F2")
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.font.size = Pt(10)

            doc.add_paragraph()  # spacer

        # ── Special Topics ──
        if newsletter.special_topics:
            h = doc.add_heading("SPECIAL TOPICS", level=1)
            h.runs[0].font.color.rgb = RGBColor(0x1E, 0x40, 0xAF)

            for topic in newsletter.special_topics:
                self._add_special_topic(doc, topic)

            doc.add_paragraph()  # spacer

        # ── Regional Sections ──
        for section_data in newsletter.sections:
            h = doc.add_heading(section_data.region_label, level=1)
            h.runs[0].font.color.rgb = RGBColor(0x1A, 0x23, 0x32)

            for item in section_data.items:
                self._add_news_item(doc, item)

            doc.add_paragraph()  # spacer

        # ── Footer / Disclaimer ──
        doc.add_paragraph()
        disclaimer = doc.add_paragraph()
        disclaimer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = disclaimer.add_run(
            "This newsletter is generated automatically for internal use only. "
            "Not investment advice. Sources include RSS feeds and NewsAPI. "
            f"Generated on {newsletter.date.isoformat()}"
        )
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
        run.font.italic = True

        logger.info("Rendered DOCX newsletter")
        return doc

    def _add_special_topic(self, doc: Document, topic: SpecialTopic) -> None:
        # Topic title
        h = doc.add_heading(topic.title, level=2)
        h.runs[0].font.color.rgb = RGBColor(0x25, 0x63, 0xEB)

        region_para = doc.add_paragraph()
        run = region_para.add_run(f"Region: {topic.region}")
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        run.font.italic = True

        # Background
        label = doc.add_paragraph()
        label_run = label.add_run("Background")
        label_run.bold = True
        label_run.font.size = Pt(11)
        label_run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
        bg = doc.add_paragraph(topic.background)
        bg.paragraph_format.space_after = Pt(8)

        # Analysis
        label = doc.add_paragraph()
        label_run = label.add_run("Analysis")
        label_run.bold = True
        label_run.font.size = Pt(11)
        label_run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
        analysis = doc.add_paragraph(topic.analysis)
        analysis.paragraph_format.space_after = Pt(8)

        # Portfolio Implications
        label = doc.add_paragraph()
        label_run = label.add_run("Portfolio Implications")
        label_run.bold = True
        label_run.font.size = Pt(11)
        label_run.font.color.rgb = RGBColor(0x1E, 0x40, 0xAF)
        impl = doc.add_paragraph(topic.portfolio_implications)
        impl.paragraph_format.space_after = Pt(12)
        for run in impl.runs:
            run.font.color.rgb = RGBColor(0x1E, 0x40, 0xAF)

    def _add_news_item(self, doc: Document, item: SummarizedItem) -> None:
        # Title with urgency badge
        title_para = doc.add_paragraph()
        title_run = title_para.add_run(item.title)
        title_run.bold = True
        title_run.font.size = Pt(11)
        title_run.font.color.rgb = RGBColor(0x1A, 0x23, 0x32)

        if item.urgency == "high":
            badge = title_para.add_run("  [HIGH]")
            badge.font.size = Pt(9)
            badge.font.color.rgb = RGBColor(0xDC, 0x26, 0x26)
            badge.bold = True

        # Summary
        summary_para = doc.add_paragraph(item.summary)
        summary_para.paragraph_format.space_after = Pt(2)
        for run in summary_para.runs:
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

        # Market implication
        if item.market_implication:
            impl_para = doc.add_paragraph()
            impl_run = impl_para.add_run(item.market_implication)
            impl_run.font.size = Pt(10)
            impl_run.font.italic = True
            impl_run.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
            impl_para.paragraph_format.space_after = Pt(2)

        # Source and time
        meta_para = doc.add_paragraph()
        pub_time = item.published_at.strftime("%H:%M UTC") if item.published_at else ""
        meta_run = meta_para.add_run(f"{item.source}  |  {pub_time}")
        meta_run.font.size = Pt(8)
        meta_run.font.color.rgb = RGBColor(0xCB, 0xD5, 0xE1)
        meta_para.paragraph_format.space_after = Pt(10)

    def save(self, doc: Document, output_dir: str = "output") -> Path:
        out_dir = Path(output_dir)
        out_dir.mkdir(parents=True, exist_ok=True)
        path = out_dir / f"Daily_Macro_Brief_{date.today().isoformat()}.docx"
        doc.save(str(path))
        logger.info("Saved DOCX to %s", path)
        return path
